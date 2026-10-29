"""
Build the augmentation-metadata annotation workbooks (task T2).

Reads the augmented schemas actually used in the experiments and samples columns
(stratified by database) for human rating of the Vietnamese alias, synonyms and
description. All tables are included for alias rating. Also writes coverage statistics
(share of columns with empty metadata), which need no human input.

Usage:
    python revision/scripts/make_t2_sheets.py \
        --aug_dir augmentation/augmented_schemas \
        --vispider data/spider_vi/vispider_dev.json \
        --vibird data/bird_vi/vibird_dev.json \
        --out revision/human_tasks/T2_metadata_check
"""

import argparse
import csv
import json
import random
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).parent))
from sheet_utils import add_list_validation, write_guide, write_header, write_lists, write_rows

SEED = 2026
N_COLUMNS = {"vispider": 80, "vibird": 120}
ANNOTATORS = ["A", "B"]

GUIDE_LINES = [
    ("Mục đích", "Đánh giá metadata tiếng Việt (tên tiếng Việt, từ đồng nghĩa, mô tả) do mô hình Qwen2.5-7B-Instruct sinh tự động "
                 "cho bảng và cột của cơ sở dữ liệu. Kết quả dùng để báo cáo chất lượng bước augmentation trong bài báo."),
    ("Cách làm", "Sheet \"Cột\": mỗi dòng là một cột của một bảng. Dựa vào tên bảng, tên cột, kiểu dữ liệu và giá trị mẫu để hiểu "
                 "cột đó chứa gì, rồi chấm metadata tiếng Việt. Sheet \"Bảng\": chấm tên tiếng Việt của từng bảng. "
                 "Chỉ điền vào các ô NỀN VÀNG."),
    ("", "Hai người chấm làm ĐỘC LẬP, không trao đổi kết quả với nhau cho đến khi cả hai nộp file."),
    ("Tên tiếng Việt (1–3)", "3 = Đúng nghĩa, là cách người dùng Việt thường gọi.\n"
                             "2 = Đúng nghĩa nhưng chưa tự nhiên hoặc quá chung/quá hẹp.\n"
                             "1 = Sai nghĩa, hoặc để trống."),
    ("Từ đồng nghĩa (1–3)", "3 = Tất cả đều phù hợp.\n"
                            "2 = Có ít nhất một từ không phù hợp hoặc trùng lặp vô ích.\n"
                            "1 = Phần lớn không phù hợp, hoặc để trống."),
    ("Mô tả (1–3)", "3 = Mô tả đúng nội dung của cột.\n"
                    "2 = Đúng một phần, mơ hồ hoặc có chi tiết bịa (ví dụ đơn vị, thang điểm không có căn cứ).\n"
                    "1 = Sai, hoặc để trống."),
    ("Có thể gây nhầm cột?", "Có = metadata này có thể khiến người hoặc mô hình nhầm cột này với một cột khác trong cùng cơ sở dữ liệu "
                             "(ví dụ hai cột khác nhau nhưng cùng tên tiếng Việt)."),
    ("Không chắc chắn", "Nếu tên cột quá khó hiểu kể cả khi xem giá trị mẫu, chấm theo khả năng tốt nhất và ghi \"không chắc\" vào cột Ghi chú."),
    ("Thời gian dự kiến", "Khoảng 1 phút mỗi cột và 15 giây mỗi bảng, tổng khoảng 4 giờ. Có thể chia nhiều buổi."),
    ("Nộp lại", "Điền họ tên và ngày ở trên, lưu file với tên giữ nguyên, gửi lại cho người phụ trách."),
]

EXAMPLE = [
    ("Bảng / Cột", "student / gpa (REAL), giá trị mẫu: 3.42, 2.85"),
    ("Metadata", "Tên: điểm trung bình | Đồng nghĩa: học lực, điểm tích lũy | Mô tả: Điểm trung bình tích lũy của sinh viên."),
    ("Tên tiếng Việt", 3),
    ("Từ đồng nghĩa", 2),
    ("Mô tả", 3),
    ("Có thể gây nhầm cột?", "Không"),
    ("Ghi chú", "\"học lực\" là xếp loại, không hẳn là điểm số"),
]


def db_ids(dev_path):
    return sorted({e["db_id"] for e in json.load(open(dev_path, encoding="utf-8"))})


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--aug_dir", required=True)
    ap.add_argument("--vispider", required=True)
    ap.add_argument("--vibird", required=True)
    ap.add_argument("--out", default="revision/human_tasks/T2_metadata_check")
    args = ap.parse_args()

    rng = random.Random(SEED)
    datasets = {"vispider": db_ids(args.vispider), "vibird": db_ids(args.vibird)}
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    col_rows, table_rows, coverage = [], [], []
    for ds, dbs in datasets.items():
        by_db = defaultdict(list)
        for db in dbs:
            path = Path(args.aug_dir) / f"{db}.json"
            if not path.exists():
                raise FileNotFoundError(f"Missing augmented schema: {path}")
            schema = json.load(open(path, encoding="utf-8"))
            n_cols = n_empty = 0
            for t in schema["tables"]:
                table_rows.append((ds, db, t["name"], t.get("alias_vi", "")))
                for c in t["columns"]:
                    n_cols += 1
                    if not c.get("alias_vi"):
                        n_empty += 1
                    by_db[db].append((ds, db, t["name"], t.get("alias_vi", ""), c))
            coverage.append((ds, db, len(schema["tables"]), n_cols, n_empty))

        # proportional allocation by #columns per DB, at least 2 per DB
        total = sum(len(v) for v in by_db.values())
        n = N_COLUMNS[ds]
        alloc = {db: max(2, round(n * len(v) / total)) for db, v in by_db.items()}
        while sum(alloc.values()) > n:
            alloc[max(alloc, key=alloc.get)] -= 1
        for db in sorted(by_db):
            col_rows.extend(rng.sample(by_db[db], min(alloc[db], len(by_db[db]))))

    with open(out / "coverage.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["dataset", "db_id", "n_tables", "n_columns", "n_columns_empty_alias"])
        w.writerows(coverage)
    with open(out / "sample_ids.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["dataset", "db_id", "table", "column"])
        for ds, db, tname, _, c in col_rows:
            w.writerow([ds, db, tname, c["name"]])

    for label in ANNOTATORS:
        wb = Workbook()
        guide = wb.active
        guide.title = "Hướng dẫn"
        write_guide(guide, "Phiếu chấm T2 — Chất lượng metadata tiếng Việt", f"Người chấm {label}", GUIDE_LINES, EXAMPLE)
        ranges = write_lists(wb.create_sheet("Danh mục"), {"Điểm": [3, 2, 1], "Có/Không": ["Có", "Không"]})

        ws = wb.create_sheet("Cột", 1)
        headers = ["STT", "Dataset", "DB", "Bảng (EN)", "Tên bảng (VI)", "Cột (EN)", "Kiểu", "Giá trị mẫu",
                   "Tên cột (VI)", "Từ đồng nghĩa (VI)", "Mô tả (VI)",
                   "Tên tiếng Việt (1–3)", "Từ đồng nghĩa (1–3)", "Mô tả (1–3)", "Có thể gây nhầm cột?",
                   "Đề xuất sửa", "Ghi chú"]
        widths = [6, 10, 18, 18, 18, 20, 9, 28, 22, 30, 40, 12, 12, 12, 12, 30, 25]
        inputs = {12, 13, 14, 15, 16, 17}
        write_header(ws, headers, widths, inputs)
        rows = [[i, ds, db, tname, talias, c["name"], c.get("type", ""), ", ".join(map(str, c.get("sample_values", []))),
                 c.get("alias_vi", ""), ", ".join(c.get("synonym", [])), c.get("description", "")] + [None] * 6
                for i, (ds, db, tname, talias, c) in enumerate(col_rows, start=1)]
        write_rows(ws, rows, inputs)
        for col in (12, 13, 14):
            add_list_validation(ws, col, len(rows), ranges["Điểm"])
        add_list_validation(ws, 15, len(rows), ranges["Có/Không"])

        ws = wb.create_sheet("Bảng", 2)
        headers = ["STT", "Dataset", "DB", "Bảng (EN)", "Tên bảng (VI)", "Tên tiếng Việt (1–3)", "Đề xuất sửa", "Ghi chú"]
        widths = [6, 10, 22, 28, 30, 12, 30, 25]
        inputs = {6, 7, 8}
        write_header(ws, headers, widths, inputs)
        rows = [[i, ds, db, tname, talias] + [None] * 3 for i, (ds, db, tname, talias) in enumerate(table_rows, start=1)]
        write_rows(ws, rows, inputs)
        add_list_validation(ws, 6, len(rows), ranges["Điểm"])

        path = out / f"T2_metadata_annotator_{label}.xlsx"
        wb.save(path)
        print(f"Saved {path}  ({len(col_rows)} columns, {len(table_rows)} tables)")


if __name__ == "__main__":
    main()
