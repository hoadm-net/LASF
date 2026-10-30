"""
Build the annotator information form (revision/human_tasks/annotator_info.xlsx).

One row per annotator and task; the person in charge fills it in. Reviewers may ask who
annotated the data, so the form records background, language proficiency and whether the
annotator took part in translating the datasets.

Usage:
    python revision/scripts/make_annotator_info.py
"""

import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from sheet_utils import add_list_validation, write_guide, write_header, write_lists, write_rows

OUT = Path("revision/human_tasks/annotator_info.xlsx")
GUIDE_LINES = [
    ("Mục đích", "Ghi thông tin người chấm để mô tả quy trình đánh giá trong bài báo (reviewer có thể hỏi). "
                 "Người phụ trách điền, mỗi người chấm một dòng cho mỗi việc."),
    ("Cách điền", "Chỉ điền vào các ô NỀN VÀNG ở sheet \"Người chấm\". Các cột có danh sách thì chọn trong danh sách."),
    ("Lưu ý", "Người chấm không nên là người đã dịch bộ dữ liệu ViSpider/ViBIRD. Nếu có, ghi rõ ở cột tương ứng."),
    ("Bảo mật", "Bài báo chỉ báo cáo thông tin tổng hợp (ví dụ: \"hai học viên cao học ngành KHMT, IELTS ≥ 7.0\"), "
                "không công bố họ tên nếu người chấm không đồng ý."),
]
EXAMPLE = [
    ("Việc", "T1"), ("Vai trò", "Người chấm A"), ("Trình độ", "Thạc sĩ"),
    ("Chuyên ngành", "Khoa học máy tính"), ("Tiếng Anh", "IELTS 7.5"),
    ("Đọc được SQL?", "Có"), ("Tham gia dịch bộ dữ liệu?", "Không"),
    ("Đồng ý ghi tên trong lời cảm ơn?", "Có"),
]


def main():
    wb = Workbook()
    guide = wb.active
    guide.title = "Hướng dẫn"
    write_guide(guide, "Thông tin người chấm (T1–T3)", "Người phụ trách", GUIDE_LINES, EXAMPLE)
    ranges = write_lists(wb.create_sheet("Danh mục"), {
        "Việc": ["T1", "T2", "T3"], "Vai trò": ["Người chấm A", "Người chấm B", "Người phân xử"],
        "Trình độ": ["Đại học", "Thạc sĩ", "Tiến sĩ", "Khác"], "Có/Không": ["Có", "Không"]})
    ws = wb.create_sheet("Người chấm", 1)
    headers = ["STT", "Việc", "Vai trò", "Họ tên", "Trình độ", "Chuyên ngành", "Đơn vị công tác/học tập",
               "Tiếng Anh (chứng chỉ hoặc mô tả)", "Đọc được SQL?", "Tham gia dịch bộ dữ liệu?",
               "Ngày giao", "Ngày nộp", "Đồng ý ghi tên trong lời cảm ơn?", "Ghi chú"]
    widths = [6, 8, 16, 24, 12, 22, 26, 24, 12, 14, 12, 12, 16, 30]
    inputs = set(range(2, 15))
    write_header(ws, headers, widths, inputs)
    rows = [[i] + [None] * 13 for i in range(1, 10)]
    write_rows(ws, rows, inputs)
    n = len(rows)
    add_list_validation(ws, 2, n, ranges["Việc"])
    add_list_validation(ws, 3, n, ranges["Vai trò"])
    add_list_validation(ws, 5, n, ranges["Trình độ"])
    for col in (9, 10, 13):
        add_list_validation(ws, col, n, ranges["Có/Không"])
    wb.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
