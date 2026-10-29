"""Shared helpers for building annotation workbooks."""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3864")
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
LOCKED_FILL = PatternFill("solid", start_color="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def write_header(ws, headers, widths, input_cols):
    for j, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEADER_FILL if j not in input_cols else PatternFill("solid", start_color="BF8F00")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "C2"


def write_rows(ws, rows, input_cols, start_row=2):
    for i, row in enumerate(rows, start=start_row):
        for j, val in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=val)
            c.font = Font(name=FONT, size=10)
            c.alignment = WRAP
            c.border = BORDER
            c.fill = INPUT_FILL if j in input_cols else LOCKED_FILL
    for j in input_cols:  # style empty input cells too
        for i in range(start_row, start_row + len(rows)):
            ws.cell(row=i, column=j).fill = INPUT_FILL


def add_list_validation(ws, col, n_rows, formula, prompt=None, start_row=2):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    dv.error = "Chọn một giá trị trong danh sách."
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    letter = get_column_letter(col)
    dv.add(f"{letter}{start_row}:{letter}{start_row + n_rows - 1}")


def write_lists(ws, lists):
    """lists: {header: [values]} written column by column; returns {header: absolute range}."""
    ranges = {}
    for j, (name, values) in enumerate(lists.items(), start=1):
        letter = get_column_letter(j)
        h = ws.cell(row=1, column=j, value=name)
        h.font = Font(name=FONT, bold=True)
        for i, v in enumerate(values, start=2):
            ws.cell(row=i, column=j, value=v).font = Font(name=FONT)
        ws.column_dimensions[letter].width = 34
        ranges[name] = f"'{ws.title}'!${letter}$2:${letter}${len(values) + 1}"
    return ranges


def write_guide(ws, title, annotator_label, lines, example=None):
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A3"] = "Người chấm"
    ws["A3"].font = Font(name=FONT, bold=True)
    ws["B3"] = annotator_label
    ws["A4"] = "Họ tên"
    ws["A4"].font = Font(name=FONT, bold=True)
    ws["B4"].fill = INPUT_FILL
    ws["A5"] = "Ngày hoàn thành"
    ws["A5"].font = Font(name=FONT, bold=True)
    ws["B5"].fill = INPUT_FILL
    r = 7
    for key, text in lines:
        a = ws.cell(row=r, column=1, value=key)
        a.font = Font(name=FONT, bold=bool(key))
        a.alignment = WRAP
        b = ws.cell(row=r, column=2, value=text)
        b.font = Font(name=FONT)
        b.alignment = WRAP
        r += 1
    if example:
        r += 1
        ws.cell(row=r, column=1, value="Ví dụ một dòng đã chấm").font = Font(name=FONT, bold=True)
        r += 1
        for key, val in example:
            ws.cell(row=r, column=1, value=key).font = Font(name=FONT, italic=True)
            c = ws.cell(row=r, column=2, value=val)
            c.font = Font(name=FONT, italic=True)
            c.alignment = WRAP
            r += 1
