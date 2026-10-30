"""
Build the T3 error-analysis workbooks (one per annotator).

120 wrong predictions, seed 2026:
  A  ViBIRD, 7B: baseline correct, full wrong (40)       -> the full prediction is reviewed
  B  ViSpider, 0.5B: full correct, baseline wrong (40)   -> the baseline prediction is reviewed
  C  wrong under all five configs, 7B (20 ViSpider + 20 ViBIRD) -> the full prediction is reviewed

Only predictions that execute are sampled: non-executable SQL is identified automatically
(A1) and needs no manual check. Each row shows the question, evidence, the schema exactly as rendered in the model's prompt,
the predicted and gold SQL, execution status and the automatic category (A3). Annotators
confirm or correct the category and flag rows usable as examples in the paper.

Usage:
    python revision/scripts/make_t3_sheets.py
"""

import argparse
import csv
import json
import random
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from sheet_utils import FONT, add_list_validation, write_guide, write_header, write_lists, write_rows
from utils.data_loader import load_spider_vi, load_bird_vi, load_schema_from_db, load_augmented_schema
from linking.schema_linker import SchemaLinker
from prompts.templates import _format_schema

SEED = 2026
ANNOTATORS = ["A", "B"]
RESULTS = Path("evaluation/results_v2")
CATEGORY_VI = {
    "invalid_sql": "SQL không chạy được",
    "gold_missing": "Thiếu bảng cần thiết trong schema",
    "wrong_tables": "Chọn sai bảng",
    "wrong_columns": "Sai cột",
    "wrong_agg_group": "Sai hàm gộp hoặc GROUP BY",
    "wrong_condition": "Sai điều kiện hoặc giá trị",
    "wrong_order": "Sai sắp xếp hoặc LIMIT",
    "other": "Khác",
}
CONFIG_FLAGS = {"baseline": (False, False, False), "full": (True, True, True)}  # aug_retrieval, aug_prompt, link
MAX_CELL = 30000

GUIDE_LINES = [
    ("Mục đích", "Kiểm tra nhãn lỗi do máy gán tự động cho các câu mô hình trả lời sai, và chọn ví dụ minh họa cho bài báo."),
    ("Cách làm", "Mỗi dòng là một câu mô hình trả lời sai. Đọc câu hỏi, evidence (nếu có), schema mô hình nhận được, "
                 "SQL mô hình sinh và SQL tham chiếu. Sau đó điền các ô màu vàng ở sheet \"Lỗi\"."),
    ("Ô cần điền", "Nhãn lỗi đúng (chọn trong danh sách) · Câu hỏi/SQL tham chiếu có vấn đề? · Dùng làm ví dụ trong bài? · Ghi chú. "
                   "Cột \"Nhãn tự động đúng?\" tự tính, không cần điền."),
    ("Chọn nhãn", "Chọn nguyên nhân chính khiến kết quả sai. Nếu có nhiều lỗi, chọn lỗi xuất hiện sớm nhất theo thứ tự danh mục dưới đây."),
    ("", ""),
    ("Danh mục nhãn", ""),
    (CATEGORY_VI["invalid_sql"], "SQL bị lỗi cú pháp, dùng bảng/cột không tồn tại, hoặc chạy quá thời gian. "
                                   "Phiếu này chỉ gồm các SQL chạy được, nên nhãn này hiếm khi dùng."),
    (CATEGORY_VI["gold_missing"], "Schema mô hình nhận được thiếu một bảng mà câu trả lời đúng cần dùng (do bước lọc schema bỏ mất)."),
    (CATEGORY_VI["wrong_tables"], "Schema có đủ bảng nhưng mô hình dùng sai bảng hoặc thiếu/thừa bảng."),
    (CATEGORY_VI["wrong_columns"], "Đúng bảng nhưng chọn sai cột (trong SELECT, JOIN hoặc điều kiện)."),
    (CATEGORY_VI["wrong_agg_group"], "Sai COUNT/SUM/AVG/MIN/MAX hoặc sai GROUP BY."),
    (CATEGORY_VI["wrong_condition"], "Sai điều kiện WHERE/HAVING: sai toán tử, sai giá trị hằng, thiếu/thừa điều kiện."),
    (CATEGORY_VI["wrong_order"], "Sai ORDER BY, chiều sắp xếp hoặc LIMIT."),
    (CATEGORY_VI["other"], "Lỗi khác: DISTINCT, điều kiện JOIN, định dạng giá trị, hoặc SQL tham chiếu có vấn đề."),
    ("", ""),
    ("Nhóm", "A: ViBIRD, model 7B — baseline đúng nhưng full sai (xem SQL của full). "
             "B: ViSpider, model 0.5B — full đúng nhưng baseline sai (xem SQL của baseline). "
             "C: sai ở cả năm config, model 7B (xem SQL của full)."),
    ("Lưu ý", "Không trao đổi với người chấm còn lại trong khi chấm. Thời gian dự kiến khoảng 3 giờ."),
]
EXAMPLE = [
    ("Câu hỏi", "Có bao nhiêu ca sĩ đến từ Pháp?"),
    ("SQL mô hình sinh", "SELECT COUNT(*) FROM singer WHERE Country = 'Pháp'"),
    ("SQL tham chiếu", "SELECT count(*) FROM singer WHERE country = 'France'"),
    ("Nhãn tự động", CATEGORY_VI["wrong_condition"]),
    ("Nhãn lỗi đúng", CATEGORY_VI["wrong_condition"]),
    ("Câu hỏi/SQL tham chiếu có vấn đề?", "Không"),
    ("Dùng làm ví dụ trong bài?", "Có"),
    ("Ghi chú", "Mô hình dịch giá trị sang tiếng Việt thay vì giữ giá trị trong database"),
]


def load_run(dataset, stem):
    with open(RESULTS / dataset / f"{stem}.jsonl", encoding="utf-8") as f:
        return {r["question_id"]: r for r in map(json.loads, f)}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cfg", default="configs/config.yaml")
    ap.add_argument("--out", default="revision/human_tasks/T3_error_analysis")
    args = ap.parse_args()
    cfg = yaml.safe_load(open(args.cfg))
    rng = random.Random(SEED)

    categories = {}
    for ds in ("spider_vi", "bird_vi"):
        with open(f"evaluation/analysis_v2/errors/{ds}_per_question.jsonl") as f:
            for r in map(json.loads, f):
                categories[(ds, r["run"], r["question_id"])] = r["category"]

    picks = []  # (group, dataset, stem, config, question_id)
    bird_b, bird_f = load_run("bird_vi", "baseline_qwen2.5-coder-7b"), load_run("bird_vi", "full_qwen2.5-coder-7b")
    pool = sorted(q for q in bird_b if bird_b[q]["execution_accuracy"] and not bird_f[q]["execution_accuracy"]
                  and bird_f[q]["exec_status"] == "ok")
    picks += [("A", "bird_vi", "full_qwen2.5-coder-7b", "full", q) for q in sorted(rng.sample(pool, 40))]

    sp_b, sp_f = load_run("spider_vi", "baseline_qwen2.5-coder-0.5b"), load_run("spider_vi", "full_qwen2.5-coder-0.5b")
    pool = sorted(q for q in sp_b if sp_f[q]["execution_accuracy"] and not sp_b[q]["execution_accuracy"]
                  and sp_b[q]["exec_status"] == "ok")
    picks += [("B", "spider_vi", "baseline_qwen2.5-coder-0.5b", "baseline", q) for q in sorted(rng.sample(pool, 40))]

    for ds in ("spider_vi", "bird_vi"):
        runs = [load_run(ds, f"{c}_qwen2.5-coder-7b") for c in ("baseline", "plusaug", "pluslink", "augret", "full")]
        pool = sorted(q for q in runs[0] if not any(run[q]["execution_accuracy"] for run in runs)
                      and runs[-1][q]["exec_status"] == "ok")
        picks += [("C", ds, "full_qwen2.5-coder-7b", "full", q) for q in sorted(rng.sample(pool, 20))]

    # Rebuild the schema each model saw (top-10 e5 retrieval on augmented columns for `full`)
    linker = SchemaLinker(model_name=cfg["linking"]["model"], device="cpu", top_k=cfg["linking"]["top_k"])
    examples = {}
    for ds, loader in (("spider_vi", load_spider_vi), ("bird_vi", load_bird_vi)):
        for e in loader(cfg["data"][ds]["dev"], lang="vi"):
            examples[(ds, e.question_id)] = e
    run_cache, rows, ids, n_mismatch = {}, [], [], 0
    for i, (group, ds, stem, config, qid) in enumerate(picks, start=1):
        if (ds, stem) not in run_cache:
            run_cache[(ds, stem)] = load_run(ds, stem)
        rec, ex = run_cache[(ds, stem)][qid], examples[(ds, qid)]
        aug_retrieval, aug_prompt, use_link = CONFIG_FLAGS[config]
        if aug_retrieval or aug_prompt:
            schema = load_augmented_schema(cfg["augmentation"]["output_dir"], ex.db_id)
        else:
            schema = load_schema_from_db(str(Path(cfg["data"][ds]["db_dir"]) / ex.db_id / f"{ex.db_id}.sqlite"), ex.db_id)
        if use_link:
            schema = linker.retrieve(ex.question, schema)
            if {t.name for t in schema.tables} != set(rec["tables_kept"]):
                n_mismatch += 1
        schema_text = _format_schema(schema, augmented=aug_prompt)
        if len(schema_text) > MAX_CELL:
            schema_text = schema_text[:MAX_CELL] + "\n... (cắt bớt để vừa ô Excel)"
        category = categories.get((ds, stem, qid), "other")
        model = "Qwen2.5-Coder-7B" if "7b" in stem else "Qwen2.5-Coder-0.5B"
        rows.append([i, group, "ViSpider" if ds == "spider_vi" else "ViBIRD", f"{model} / {config}", ex.question,
                     ex.evidence or "", schema_text, rec["pred_sql"], ex.gold_sql, rec["exec_status"],
                     CATEGORY_VI[category], None, None, None, None, None])
        ids.append([i, group, ds, stem, qid, ex.db_id, category])

    print(f"Rebuilt filtered schemas whose tables differ from the run (CPU vs GPU encoding): {n_mismatch}")
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    with open(out / "sample_ids.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["STT", "group", "dataset", "run", "question_id", "db_id", "auto_category"])
        w.writerows(ids)

    n = len(rows)
    for label in ANNOTATORS:
        wb = Workbook()
        guide = wb.active
        guide.title = "Hướng dẫn"
        write_guide(guide, "Phiếu chấm T3 — Phân tích lỗi", f"Người chấm {label}", GUIDE_LINES, EXAMPLE)
        ranges = write_lists(wb.create_sheet("Danh mục"),
                             {"Nhãn lỗi": list(CATEGORY_VI.values()), "Có/Không": ["Có", "Không"]})

        ws = wb.create_sheet("Lỗi", 1)
        headers = ["STT", "Nhóm", "Dataset", "Model / config", "Câu hỏi", "Evidence", "Schema mô hình nhận được",
                   "SQL mô hình sinh", "SQL tham chiếu", "Trạng thái thực thi", "Nhãn tự động",
                   "Nhãn lỗi đúng", "Nhãn tự động đúng?", "Câu hỏi/SQL tham chiếu có vấn đề?",
                   "Dùng làm ví dụ trong bài?", "Ghi chú"]
        widths = [6, 7, 10, 20, 40, 30, 60, 45, 45, 12, 22, 24, 12, 14, 12, 30]
        inputs = {12, 14, 15, 16}
        write_header(ws, headers, widths, inputs)
        write_rows(ws, rows, inputs)
        for r in range(2, n + 2):
            ws.cell(row=r, column=13, value=f'=IF(L{r}="","",IF(L{r}=K{r},"Có","Không"))')
            ws.row_dimensions[r].height = 150
        add_list_validation(ws, 12, n, ranges["Nhãn lỗi"])
        add_list_validation(ws, 14, n, ranges["Có/Không"])
        add_list_validation(ws, 15, n, ranges["Có/Không"])

        summary = wb.create_sheet("Tổng hợp", 2)
        summary.column_dimensions["A"].width = 40
        summary.column_dimensions["B"].width = 16
        summary.column_dimensions["C"].width = 16
        stats = [
            ("Số dòng đã chấm", f"=COUNTA('Lỗi'!L2:L{n + 1})"),
            ("Tỉ lệ nhãn tự động đúng", f"=IF(B2=0,\"\",COUNTIF('Lỗi'!M2:M{n + 1},\"Có\")/B2)"),
            ("Số câu đề xuất làm ví dụ", f"=COUNTIF('Lỗi'!O2:O{n + 1},\"Có\")"),
            ("Số câu hỏi/SQL tham chiếu có vấn đề", f"=COUNTIF('Lỗi'!N2:N{n + 1},\"Có\")"),
        ]
        for r, (name, formula) in enumerate(stats, start=1):
            summary.cell(row=r + 1, column=1, value=name).font = Font(name=FONT)
            summary.cell(row=r + 1, column=2, value=formula).font = Font(name=FONT)
        summary.cell(row=1, column=1, value="Chỉ số").font = Font(name=FONT, bold=True)
        summary.cell(row=1, column=2, value="Giá trị").font = Font(name=FONT, bold=True)
        summary["B3"].number_format = "0.0%"
        r0 = len(stats) + 3
        for j, h in enumerate(["Nhãn", "Nhãn tự động", "Nhãn người chấm"], start=1):
            summary.cell(row=r0, column=j, value=h).font = Font(name=FONT, bold=True)
        for k, name in enumerate(CATEGORY_VI.values(), start=r0 + 1):
            summary.cell(row=k, column=1, value=name).font = Font(name=FONT)
            summary.cell(row=k, column=2, value=f"=COUNTIF('Lỗi'!K2:K{n + 1},A{k})").font = Font(name=FONT)
            summary.cell(row=k, column=3, value=f"=COUNTIF('Lỗi'!L2:L{n + 1},A{k})").font = Font(name=FONT)

        wb.calculation.fullCalcOnLoad = True
        path = out / f"T3_error_annotator_{label}.xlsx"
        wb.save(path)
        print(f"Saved {path}  ({n} rows)")


if __name__ == "__main__":
    main()
